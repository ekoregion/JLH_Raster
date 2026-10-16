#!/usr/bin/env python3
"""
Kalkulasi 6 Indikator Jasa Lingkungan Hidup (JLH) Sumatera Barat 2025
Menggunakan Multiscale Land Cover 2025 (0.5 meter), KBA (0.5m), dan KVA (0.5m).
Output: Multi-band GeoTIFF terkompresi DEFLATE (predictor 3) per-grid + Virtual Raster (.vrt) mozaik.
Dijalankan secara multi-core paralel (16 worker).
"""

import os
import glob
import time
import argparse
import multiprocessing as mp
import xml.etree.ElementTree as ET
import numpy as np
import rasterio
from rasterio.windows import from_bounds
import openpyxl

# Konfigurasi JLH Sumatera Barat
JLH_CONFIG = {
    'pga': {'name': 'Pengaturan Air',             'w_kba': 0.28, 'w_kva': 0.12, 'w_pl': 0.60},
    'phk': {'name': 'Pendukung Habitat & Kehati', 'w_kba': 0.24, 'w_kva': 0.16, 'w_pl': 0.60},
    'ppk': {'name': 'Penyerap & Simpan Karbon',   'w_kba': 0.20, 'w_kva': 0.20, 'w_pl': 0.60},
    'pya': {'name': 'Penyedia Air',               'w_kba': 0.32, 'w_kva': 0.08, 'w_pl': 0.60},
    'pgn': {'name': 'Penyedia Pangan (Sumatera)', 'w_kba': 0.28, 'w_kva': 0.12, 'w_pl': 0.60},
    'pku': {'name': 'Pengatur Kualitas Udara',    'w_kba': 0.15, 'w_kva': 0.25, 'w_pl': 0.60}
}
JLH_KEYS = list(JLH_CONFIG.keys())

PL_LOOKUP = {
    1:  {'pga': 5, 'phk': 1, 'ppk': 1, 'pya': 5, 'pgn': 4, 'pku': 5},
    2:  {'pga': 5, 'phk': 5, 'ppk': 5, 'pya': 3, 'pgn': 4, 'pku': 5},
    4:  {'pga': 4, 'phk': 3, 'ppk': 3, 'pya': 3, 'pgn': 4, 'pku': 5},
    5:  {'pga': 3, 'phk': 1, 'ppk': 1, 'pya': 3, 'pgn': 4, 'pku': 3},
    7:  {'pga': 2, 'phk': 1, 'ppk': 1, 'pya': 2, 'pgn': 1, 'pku': 2},
    8:  {'pga': 1, 'phk': 1, 'ppk': 1, 'pya': 1, 'pgn': 1, 'pku': 1},
    11: {'pga': 3, 'phk': 2, 'ppk': 2, 'pya': 1, 'pgn': 2, 'pku': 3},
    12: {'pga': 3, 'phk': 2, 'ppk': 2, 'pya': 2, 'pgn': 3, 'pku': 4},
    13: {'pga': 3, 'phk': 3, 'ppk': 2, 'pya': 2, 'pgn': 3, 'pku': 3},
    14: {'pga': 1, 'phk': 1, 'ppk': 1, 'pya': 1, 'pgn': 1, 'pku': 1},
    15: {'pga': 1, 'phk': 1, 'ppk': 1, 'pya': 2, 'pgn': 1, 'pku': 1}
}

# Variabel global worker
worker_src_pl = None
shared_remap = None

def init_worker(pl_raster_path, remap_dict):
    global worker_src_pl, shared_remap
    worker_src_pl = rasterio.open(pl_raster_path)
    shared_remap = remap_dict

def load_scoring_tables(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_kba = wb['Skor KBA']
    ws_kva = wb['Skor KVA']

    pl_remap = {k: np.zeros(256, dtype=np.float32) for k in JLH_KEYS}
    for c_val, scores in PL_LOOKUP.items():
        for k in JLH_KEYS:
            pl_remap[k][c_val] = scores[k]

    kba_remap = {k: np.zeros(256, dtype=np.float32) for k in JLH_KEYS}
    for r in list(ws_kba.iter_rows(values_only=True))[2:]:
        if r[0] is not None and str(r[0]).strip().isdigit():
            val = int(r[0])
            for idx, k in enumerate(JLH_KEYS):
                if r[2+idx] is not None:
                    kba_remap[k][val] = float(r[2+idx])

    kva_remap = {k: np.zeros(256, dtype=np.float32) for k in JLH_KEYS}
    for r in list(ws_kva.iter_rows(values_only=True))[2:]:
        if r[0] is not None and str(r[0]).strip().isdigit():
            val = int(r[0])
            for idx, k in enumerate(JLH_KEYS):
                if r[2+idx] is not None:
                    kva_remap[k][val] = float(r[2+idx])

    return {'pl': pl_remap, 'kba': kba_remap, 'kva': kva_remap}

def process_tile_task(task_args):
    kba_path, kva_path, out_path = task_args
    if os.path.exists(out_path):
        return 'skip'

    global worker_src_pl, shared_remap
    pl_remap = shared_remap['pl']
    kba_remap = shared_remap['kba']
    kva_remap = shared_remap['kva']

    with rasterio.open(kba_path) as src_kba, rasterio.open(kva_path) as src_kva:
        kba_data = src_kba.read(1)
        kva_data = src_kva.read(1)
        bounds = src_kba.bounds
        h, w = src_kba.height, src_kba.width
        prof = src_kba.profile.copy()

    # Windowed read dari LULC 2025
    win_pl = from_bounds(*bounds, transform=worker_src_pl.transform)
    pl_data = worker_src_pl.read(1, window=win_pl, out_shape=(h, w), boundless=True, fill_value=0)

    nodata_mask = (pl_data == 0) | (pl_data == 255)

    # Cek jika seluruh tile adalah nodata/laut
    if np.all(nodata_mask) and np.all(kba_data == 0) and np.all(kva_data == 0):
        prof.update({
            'count': 6,
            'dtype': 'float32',
            'nodata': 0.0,
            'compress': 'deflate',
            'predictor': 3,
            'tiled': True,
            'blockxsize': 512,
            'blockysize': 512
        })
        empty = np.zeros((6, h, w), dtype=np.float32)
        with rasterio.open(out_path, 'w', **prof) as dst:
            dst.write(empty)
        return 'empty'

    # Kalkulasi 6 Band JLH
    jlh_result = np.zeros((6, h, w), dtype=np.float32)
    for idx, k in enumerate(JLH_KEYS):
        cfg = JLH_CONFIG[k]
        sk_kba = kba_remap[k][kba_data]
        sk_kva = kva_remap[k][kva_data]
        sk_pl = pl_remap[k][pl_data]

        band = (sk_kba * cfg['w_kba']) + (sk_kva * cfg['w_kva']) + (sk_pl * cfg['w_pl'])
        band[nodata_mask] = 0.0
        jlh_result[idx] = band

    prof.update({
        'count': 6,
        'dtype': 'float32',
        'nodata': 0.0,
        'compress': 'deflate',
        'predictor': 3,
        'tiled': True,
        'blockxsize': 512,
        'blockysize': 512
    })

    with rasterio.open(out_path, 'w', **prof) as dst:
        dst.write(jlh_result)
        for b_idx, k in enumerate(JLH_KEYS, 1):
            dst.set_band_description(b_idx, f"JLH_{k.upper()}_{JLH_CONFIG[k]['name']}")

    return 'processed'

def create_indicator_vrt(tif_list, output_vrt_path, band_index=1):
    if not tif_list:
        return
    
    meta_list = []
    min_x, max_x = float('inf'), float('-inf')
    min_y, max_y = float('inf'), float('-inf')
    crs_wkt = None
    res_x, res_y = 0.5, 0.5

    for p in tif_list:
        with rasterio.open(p) as src:
            b = src.bounds
            min_x = min(min_x, b.left)
            max_x = max(max_x, b.right)
            min_y = min(min_y, b.bottom)
            max_y = max(max_y, b.top)
            if crs_wkt is None:
                crs_wkt = src.crs.to_wkt()
                res_x = src.transform.a
                res_y = abs(src.transform.e)
            meta_list.append({
                'rel_path': os.path.relpath(p, os.path.dirname(output_vrt_path)),
                'bounds': b,
                'width': src.width,
                'height': src.height
            })
            
    vrt_width = int(round((max_x - min_x) / res_x))
    vrt_height = int(round((max_y - min_y) / res_y))
    geotransform = f'{min_x}, {res_x}, 0.0, {max_y}, 0.0, {-res_y}'

    root = ET.Element('VRTDataset', rasterXSize=str(vrt_width), rasterYSize=str(vrt_height))
    srs = ET.SubElement(root, 'SRS', dataAxisToSRSAxisMapping='1,2')
    srs.text = crs_wkt
    gt = ET.SubElement(root, 'GeoTransform')
    gt.text = geotransform

    band = ET.SubElement(root, 'VRTRasterBand', dataType='Float32', band='1')
    nodata = ET.SubElement(band, 'NoDataValue')
    nodata.text = '0.0'
    color = ET.SubElement(band, 'ColorInterp')
    color.text = 'Gray'

    for m in meta_list:
        dst_x = int(round((m['bounds'].left - min_x) / res_x))
        dst_y = int(round((max_y - m['bounds'].top) / res_y))
        w, h = m['width'], m['height']

        cs = ET.SubElement(band, 'ComplexSource')
        sf = ET.SubElement(cs, 'SourceFilename', relativeToVRT='1')
        sf.text = m['rel_path']
        sb = ET.SubElement(cs, 'SourceBand')
        sb.text = str(band_index)
        sp = ET.SubElement(cs, 'SourceProperties', RasterXSize=str(w), RasterYSize=str(h), DataType='Float32', BlockXSize='512', BlockYSize='512')
        sr = ET.SubElement(cs, 'SrcRect', xOff='0', yOff='0', xSize=str(w), ySize=str(h))
        dr = ET.SubElement(cs, 'DstRect', xOff=str(dst_x), yOff=str(dst_y), xSize=str(w), ySize=str(h))
        nd = ET.SubElement(cs, 'NODATA')
        nd.text = '0.0'

    tree = ET.ElementTree(root)
    tree.write(output_vrt_path, encoding='utf-8', xml_declaration=False)
    print(f"🗺️  VRT Terbuat: {output_vrt_path}")

def main():
    parser = argparse.ArgumentParser(description="Kalkulasi JLH Sumatera Barat 2025")
    parser.add_argument('--workers', type=int, default=16, help="Jumlah parallel worker processes")
    parser.add_argument('--limit', type=int, default=None, help="Batas tile untuk testing")
    args = parser.parse_args()

    t_start = time.time()
    print("=" * 65)
    print("🌿 PEMODELAN JASA LINGKUNGAN HIDUP (JLH) SUMATERA BARAT 2025")
    print("=" * 65)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    kba_dir = os.path.join(base_dir, 'data', 'kba')
    kva_dir = os.path.join(base_dir, 'data', 'kva')
    pl_path = os.path.join(base_dir, 'data', 'penutupan_lahan', 'multiscale_2025_seamless_class.tif')
    excel_path = os.path.join(base_dir, 'data', 'Skor_JLH_2026.xlsx')
    out_grid_dir = os.path.join(base_dir, 'output', 'jlh_sumbar_2025', 'grids')
    out_vrt_dir = os.path.join(base_dir, 'output', 'jlh_sumbar_2025')

    os.makedirs(out_grid_dir, exist_ok=True)

    print("📖 Membaca tabel skor dan parameter bobot...")
    remap_dict = load_scoring_tables(excel_path)
    print("  • 6 Indikator JLH: PGA, PHK, PPK, PYA, PGN, PKU terkonfigurasi!")

    kba_files = sorted(glob.glob(os.path.join(kba_dir, '*.tif')))
    print(f"\n📂 Menyiapkan {len(kba_files):,} task grid...")

    tasks = []
    for kba_f in kba_files:
        bname = os.path.basename(kba_f)
        kva_f = os.path.join(kva_dir, bname.replace('_kba.tif', '_kva.tif'))
        out_f = os.path.join(out_grid_dir, bname.replace('_kba.tif', '_JLH_6band.tif'))
        if os.path.exists(kva_f):
            tasks.append((kba_f, kva_f, out_f))

    if args.limit:
        tasks = tasks[:args.limit]

    print(f"  • Total grid yang valid: {len(tasks):,} file")
    print(f"🚀 Memulai komputasi paralel dengan {args.workers} worker...")

    t0 = time.time()
    processed_cnt = 0
    skipped_cnt = 0

    with mp.Pool(processes=args.workers, initializer=init_worker, initargs=(pl_path, remap_dict)) as pool:
        for i, res in enumerate(pool.imap_unordered(process_tile_task, tasks, chunksize=8), 1):
            if res == 'skip':
                skipped_cnt += 1
            else:
                processed_cnt += 1

            if i % 100 == 0 or i == len(tasks):
                elapsed = time.time() - t0
                spd = i / elapsed if elapsed > 0 else 0
                rem = (len(tasks) - i) / spd if spd > 0 else 0
                print(f"  [{i:,}/{len(tasks):,}] Grid selesai ({elapsed:.1f}s, kecepatan: {spd:.1f} grid/s, sisa: {rem/60:.1f} mnt)")

    print(f"\n✅ Komputasi grid selesai dalam {time.time() - t0:.1f} detik!")
    print(f"  • Diproses : {processed_cnt:,} grid")
    print(f"  • Dilewati : {skipped_cnt:,} grid")

    print("\n📦 Membangun Virtual Raster (.vrt) untuk 6 Indikator...")
    output_tiles = sorted(glob.glob(os.path.join(out_grid_dir, '*_JLH_6band.tif')))
    
    for idx, k in enumerate(JLH_KEYS, 1):
        vrt_out = os.path.join(out_vrt_dir, f"sumbar_2025_jlh_{k}.vrt")
        create_indicator_vrt(output_tiles, vrt_out, band_index=idx)

    # Hitung total ukuran disk
    total_size = sum(os.path.getsize(f) for f in output_tiles) / (1024 * 1024)
    print("\n" + "=" * 65)
    print("🎉 SEMUA PROSES PEMODELAN JLH SUMBAR 2025 SELESAI!")
    print(f"  • Total Ukuran Grid di Disk: {total_size:.1f} MB ({total_size/1024:.2f} GB)")
    print(f"  • Folder Output Grid       : {out_grid_dir}")
    print(f"  • File Mozaik VRT          : {out_vrt_dir}/sumbar_2025_jlh_*.vrt")
    print(f"  • Waktu Total              : {(time.time() - t_start)/60:.1f} menit")
    print("=" * 65)

if __name__ == '__main__':
    main()
