import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Color Palette
    C_NAVY = RGBColor(27, 54, 93)      # #1B365D Primary Dark
    C_TEAL = RGBColor(32, 115, 140)    # #20738C Secondary Accent
    C_GOLD = RGBColor(212, 160, 23)    # #D4A017 Highlight
    C_BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
    C_DARK = RGBColor(33, 37, 41)      # #212529 Text
    C_WHITE = RGBColor(255, 255, 255)  # White
    C_GRAY = RGBColor(108, 117, 125)   # #6C757D
    C_CARD_BG = RGBColor(241, 245, 249) # #F1F5F9
    C_GREEN = RGBColor(40, 167, 69)    # #28A745 Positive
    C_RED = RGBColor(220, 53, 69)      # #DC3545 Negative
    C_HEADER_ROW = RGBColor(27, 54, 93)

    def add_header(slide, title_text, category_text="KEMENTERIAN LINGKUNGAN HIDUP / BPLH • DIREKTORAT PENYELENGGARAAN SDA BERKELANJUTAN"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        banner.fill.solid()
        banner.fill.fore_color.rgb = C_NAVY
        banner.line.color.rgb = C_NAVY
        
        # Category / Agency
        tx_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        tf_c = tx_cat.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(9.5)
        p_c.font.bold = True
        p_c.font.color.rgb = C_GOLD
        
        # Title
        tx_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.6))
        tf_t = tx_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE
        
        # Accent gold bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = C_GOLD
        bar.line.color.rgb = C_GOLD

    def add_matrix_table(slide, headers, data, left=Inches(0.8), top=Inches(1.5), width=Inches(8.0), height=Inches(5.4)):
        rows = len(data) + 1
        cols = len(headers)
        table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        tbl = table_shape.table
        
        # Format Headers
        for col_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, col_idx)
            cell.text = h_text
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_HEADER_ROW
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = C_WHITE
            
        # Format Data Rows
        for row_idx, row_vals in enumerate(data, start=1):
            is_total = (row_idx == len(data))
            bg_color = RGBColor(235, 240, 245) if is_total else (RGBColor(248, 250, 252) if row_idx % 2 == 1 else C_WHITE)
            
            for col_idx, val in enumerate(row_vals):
                cell = tbl.cell(row_idx, col_idx)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg_color
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(10)
                
                if col_idx in [0, 1]:
                    p.alignment = PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.CENTER
                else:
                    p.alignment = PP_ALIGN.RIGHT
                    
                if is_total or col_idx == 0:
                    p.font.bold = True
                
                # Colorize Deltas if column 6 or 7
                if col_idx in [6, 7] and not is_total:
                    if str(val).startswith("+"):
                        p.font.color.rgb = C_GREEN
                        p.font.bold = True
                    elif str(val).startswith("-"):
                        p.font.color.rgb = C_RED
                        p.font.bold = True
                    else:
                        p.font.color.rgb = C_DARK
                elif is_total:
                    p.font.color.rgb = C_NAVY
                else:
                    p.font.color.rgb = C_DARK
                    
        return table_shape

    def add_insight_box(slide, title, bullets, left=Inches(9.1), top=Inches(1.5), width=Inches(3.4), height=Inches(5.4)):
        # Card background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = RGBColor(203, 213, 225)
        
        # Header inside card
        c_hdr = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.65))
        c_hdr.fill.solid()
        c_hdr.fill.fore_color.rgb = C_TEAL
        c_hdr.line.color.rgb = C_TEAL
        tf_h = c_hdr.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = title.upper()
        p_h.font.size = Pt(11.5)
        p_h.font.bold = True
        p_h.font.color.rgb = C_WHITE
        p_h.alignment = PP_ALIGN.CENTER
        
        # Bullets
        tx = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.75), width - Inches(0.3), height - Inches(0.85))
        tf = tx.text_frame
        tf.word_wrap = True
        
        for i, (b_title, b_desc) in enumerate(bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.space_after = Pt(8)
            p.line_spacing = 1.15
            
            run_t = p.add_run()
            run_t.text = f"• {b_title}: "
            run_t.font.bold = True
            run_t.font.size = Pt(10)
            run_t.font.color.rgb = C_NAVY
            
            run_d = p.add_run()
            run_d.text = b_desc
            run_d.font.size = Pt(9.5)
            run_d.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 1: COVER
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY
    bg1.line.color.rgb = C_NAVY
    
    # Gold accent frame
    f_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(0.15), Inches(5.5))
    f_bar.fill.solid()
    f_bar.fill.fore_color.rgb = C_GOLD
    f_bar.line.color.rgb = C_GOLD
    
    tx_cov = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(5.0))
    tf_cov = tx_cov.text_frame
    tf_cov.word_wrap = True
    
    p0 = tf_cov.paragraphs[0]
    p0.text = "KEMENTERIAN LINGKUNGAN HIDUP / BADAN PENGENDALIAN LINGKUNGAN HIDUP (BPLH)\nDIREKTORAT PENYELENGGARAAN SUMBER DAYA ALAM BERKELANJUTAN"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD
    p0.space_after = Pt(20)
    
    p1 = tf_cov.add_paragraph()
    p1.text = "MATRIKS PERBANDINGAN LUASAN\nJASA LINGKUNGAN HIDUP (JLH) 2025"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_after = Pt(14)
    
    p2 = tf_cov.add_paragraph()
    p2.text = "Komparasi Saintifik & Spasial: Pemodelan Raster 0.5 Meter (Citra Satelit 2025)\nversus Basis Data Vektor Geodatabase Resmi (JLH_2025_ALL.gdb) di Provinsi Sumatera Barat"
    p2.font.size = Pt(14)
    p2.font.italic = True
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(28)
    
    p3 = tf_cov.add_paragraph()
    p3.text = "Cakupan: 6 Indikator JLH (PGA, PYA, PGN, PKU, PHK, PPK) • Resolusi: 0.5m & Skala Makro • Tanggal: Oktober 2026"
    p3.font.size = Pt(11)
    p3.font.color.rgb = C_GOLD

    # =========================================================================
    # SLIDE 2: RINGKASAN EKSEKUTIF
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Ringkasan Eksekutif Komparasi JLH Sumatera Barat")
    
    # 3 Summary Cards
    cards_data = [
        ("TOTAL CAKUPAN LUAS", [
            ("Raster 0.5m", "4.380.394 Ha (~175,2 Miliar piksel valid)."),
            ("Vektor GDB", "4.213.812 Ha (37.001 fitur poligon)."),
            ("Selisih (+3,95%)", "Raster menangkap ribuan pulau kecil/atol Mentawai & sempadan pantai mikro yang tidak terdeliniasi pada skala 1:50.000.")
        ], Inches(0.8), Inches(1.5), Inches(3.7), Inches(5.4)),
        ("KETAJAMAN KANOPI POHON", [
            ("Sensitivitas 50cm", "Piksel 0.5m mendeteksi tajuk kanopi pohon aktual di kebun rakyat & agroforestri."),
            ("Lonjakan Sangat Tinggi", "Pada PKU (Udara) dan PPK (Karbon), porsi Sangat Tinggi mencapai 57-65% (vs 28-46% di GDB)."),
            ("Generalisasi Poligon", "Vektor GDB meratakan mozaik pohon soliter ke kelas agregat sedang/rendah.")
        ], Inches(4.8), Inches(1.5), Inches(3.7), Inches(5.4)),
        ("PANGAN & AIR ALAMIAH", [
            ("PGN (Pangan)", "Sawah irigasi di lembah vulkanik subur terdeteksi murni pada kelas Tinggi (72,86% vs 38,97% di GDB)."),
            ("PYA (Air Bersih)", "Kelas Sedang mendominasi 71,08% pada raster karena aliran sungai mikro terpetakan utuh."),
            ("Efek Legal KK", "Penerapan Kriteria Khusus pada PHK melonjakkan zona konservasi inti dari 5,12% ke 37,48%.")
        ], Inches(8.8), Inches(1.5), Inches(3.7), Inches(5.4))
    ]
    for c_title, c_bullets, c_l, c_t, c_w, c_h in cards_data:
        add_insight_box(slide2, c_title, c_bullets, left=c_l, top=c_t, width=c_w, height=c_h)

    # =========================================================================
    # SLIDE 3: MATRIKS KARAKTERISTIK TEKNIS
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Matriks Perbandingan Karakteristik Teknis & Metodologis")
    
    tech_headers = ["Parameter / Dimensi", "Raster 0.5 Meter Seamless (2025)", "Vektor GDB Resmi (JLH_2025_ALL.gdb)"]
    tech_data = [
        ["Format & Arsitektur Data", "Cloud Optimized GeoTIFF (COG), LERC_ZSTD", "Esri File Geodatabase (MultiPolygon Feature Class)"],
        ["Resolusi Spasial & Skala", "0.5 meter x 0.5 meter (Setara Skala Tapak 1:1.000)", "Poligon Vektor Makro (Setara Skala 1:50.000)"],
        ["Unit Analisis Spasial", "Sel Piksel Kontinu Float32", "Agregasi Poligon Batas Tutupan Lahan"],
        ["Jumlah Elemen Geospasial", "~175,2 Miliar Piksel Valid se-Sumbar", "37.001 Fitur Poligon"],
        ["Total Luas Terpetakan", "4.380.394,17 Hektar (Detail pulau karang/Mentawai)", "4.213.812,00 Hektar (Batas daratan standar)"],
        ["Deteksi Kanopi & Sempadan", "Sangat Tajam: Memisahkan tajuk pohon, sempadan 5m, jalan", "Terglobalkan: Meratakan mosaik tajuk ke kelas dominan"],
        ["Kecepatan Analisis Spasial", "Sangat cepat untuk aljabar peta MCE & hidrologi mikro", "Sangat cepat untuk query agregat tabel administrasi"],
        ["Rekomendasi Penggunaan", "Skala Tapak: AMDAL detail, perizinan, relokasi, D3TLH", "Skala Makro: Neraca SDA provinsi, RPPLH, pelaporan"]
    ]
    tbl_tech = add_matrix_table(slide3, tech_headers, tech_data, left=Inches(0.8), top=Inches(1.5), width=Inches(11.7), height=Inches(5.4))
    # Adjust column widths for tech table
    tbl_tech.table.columns[0].width = Inches(2.7)
    tbl_tech.table.columns[1].width = Inches(4.5)
    tbl_tech.table.columns[2].width = Inches(4.5)

    # Helper for 6 indicator slides
    common_headers = ["Kelas JLH", "Skor", "Raster 0.5m (Ha)", "% Rast", "GDB Vektor (Ha)", "% GDB", "Selisih Luas", "Selisih %"]

    # =========================================================================
    # SLIDE 4: PGA (PENGATURAN AIR)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Matriks Perbandingan: PGA (Pengaturan Air / Water Regulation)")
    pga_data = [
        ["Sangat Rendah", "1.0–1.8", "29.073 Ha", "0,66%", "8.993 Ha", "0,21%", "+20.080 Ha", "+0,45%"],
        ["Rendah", "1.8–2.6", "357.572 Ha", "8,16%", "309.061 Ha", "7,33%", "+48.511 Ha", "+0,83%"],
        ["Sedang", "2.6–3.4", "1.149.438 Ha", "26,24%", "1.751.883 Ha", "41,57%", "-602.445 Ha", "-15,33%"],
        ["Tinggi", "3.4–4.2", "733.377 Ha", "16,74%", "716.062 Ha", "16,99%", "+17.315 Ha", "-0,25%"],
        ["Sangat Tinggi", "4.2–5.0", "2.110.935 Ha", "48,19%", "1.427.814 Ha", "33,88%", "+683.121 Ha", "+14,31%"],
        ["TOTAL DARATAN", "-", "4.380.394 Ha", "100,00%", "4.213.812 Ha", "100,00%", "+166.582 Ha", "-"]
    ]
    add_matrix_table(slide4, common_headers, pga_data)
    add_insight_box(slide4, "Analisis Spasial PGA", [
        ("Dominasi Sangat Tinggi", "Raster 0.5m menghasilkan kelas Sangat Tinggi 48,19% (+14,31% vs GDB) karena zona resapan air pegunungan vulkanik dan karst Bukit Barisan terpetakan sangat utuh."),
        ("Penciutan Kelas Sedang", "Kelas Sedang di raster hanya 26,24% (vs 41,57% di GDB) karena area vegetasi lereng curam yang sebelumnya tergeneralisasi sedang, naik ke kelas tinggi/sangat tinggi."),
        ("Akurasi Sempadan Air", "Lahan terbangun rapat di dataran banjir terpetakan jelas pada kelas Sangat Rendah (+20.080 Ha).")
    ])

    # =========================================================================
    # SLIDE 5: PYA (PENYEDIA AIR BERSIH)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Matriks Perbandingan: PYA (Penyedia Air Bersih / Water Provisioning)")
    pya_data = [
        ["Sangat Rendah", "1.0–1.8", "92.956 Ha", "2,12%", "46.574 Ha", "1,11%", "+46.382 Ha", "+1,01%"],
        ["Rendah", "1.8–2.6", "891.305 Ha", "20,35%", "2.364.620 Ha", "56,12%", "-1.473.315 Ha", "-35,77%"],
        ["Sedang", "2.6–3.4", "3.113.503 Ha", "71,08%", "1.749.860 Ha", "41,53%", "+1.363.643 Ha", "+29,55%"],
        ["Tinggi", "3.4–4.2", "240.258 Ha", "5,48%", "17.717 Ha", "0,42%", "+222.541 Ha", "+5,06%"],
        ["Sangat Tinggi", "4.2–5.0", "42.372 Ha", "0,97%", "35.041 Ha", "0,83%", "+7.331 Ha", "+0,14%"],
        ["TOTAL DARATAN", "-", "4.380.394 Ha", "100,00%", "4.213.812 Ha", "100,00%", "+166.582 Ha", "-"]
    ]
    add_matrix_table(slide5, common_headers, pya_data)
    add_insight_box(slide5, "Analisis Spasial PYA", [
        ("Pergeseran Rendah ke Sedang", "Pada GDB, kelas Rendah mencapai 56,12%. Di raster 0.5m, kelas ini berkurang drastis menjadi 20,35%, dan beralih ke kelas Sedang (71,08%)."),
        ("Peran Akuifer Alami", "KBA geomorfologi cekungan air tanah dan tutupan perkebunan/hutan sekunder dinilai lebih realistis pada skor 2,6-3,4 pada raster."),
        ("Penyedia Air Alami", "Kelas Sangat Tinggi merefleksikan badan air tawar (Danau Singkarak, Maninjau, Danau Kembar, sungai besar) seluas ~42.372 Ha.")
    ])

    # =========================================================================
    # SLIDE 6: PGN (PENYEDIA PANGAN)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Matriks Perbandingan: PGN (Penyedia Pangan / Food Provisioning)")
    pgn_data = [
        ["Sangat Rendah", "1.0–1.8", "38.670 Ha", "0,88%", "26.376 Ha", "0,63%", "+12.294 Ha", "+0,25%"],
        ["Rendah", "1.8–2.6", "307.563 Ha", "7,02%", "113.411 Ha", "2,69%", "+194.152 Ha", "+4,33%"],
        ["Sedang", "2.6–3.4", "671.672 Ha", "15,33%", "2.200.546 Ha", "52,22%", "-1.528.874 Ha", "-36,89%"],
        ["Tinggi", "3.4–4.2", "3.191.533 Ha", "72,86%", "1.641.998 Ha", "38,97%", "+1.549.535 Ha", "+33,89%"],
        ["Sangat Tinggi", "4.2–5.0", "170.956 Ha", "3,90%", "231.482 Ha", "5,49%", "-60.526 Ha", "-1,59%"],
        ["TOTAL DARATAN", "-", "4.380.394 Ha", "100,00%", "4.213.812 Ha", "100,00%", "+166.582 Ha", "-"]
    ]
    add_matrix_table(slide6, common_headers, pgn_data)
    add_insight_box(slide6, "Analisis Spasial PGN", [
        ("Dominasi Kelas Tinggi (72,86%)", "Pada raster 0.5m, kawasan pertanian pangan (sawah irigasi) dan perkebunan rakyat di tanah vulkanik subur terpetakan murni pada kelas Tinggi (skor 3.4-4.2)."),
        ("Perbedaan dengan GDB", "Di GDB, kawasan budidaya terbagi antara Sedang (52,22%) dan Tinggi (38,97%) akibat generalisasi poligon campuran."),
        ("Lumbung Pangan Sumbar", "Total area pangan potensial (Tinggi + Sangat Tinggi) mencapai 3,36 juta Ha di raster 0.5m.")
    ])

    # =========================================================================
    # SLIDE 7: PKU (KUALITAS UDARA)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Matriks Perbandingan: PKU (Pengatur Kualitas Udara / Air Quality)")
    pku_data = [
        ["Sangat Rendah", "1.0–1.8", "6.840 Ha", "0,16%", "0 Ha", "0,00%", "+6.840 Ha", "+0,16%"],
        ["Rendah", "1.8–2.6", "44.263 Ha", "1,01%", "22.878 Ha", "0,54%", "+21.385 Ha", "+0,47%"],
        ["Sedang", "2.6–3.4", "519.845 Ha", "11,87%", "1.627.534 Ha", "38,62%", "-1.107.689 Ha", "-26,75%"],
        ["Tinggi", "3.4–4.2", "938.705 Ha", "21,43%", "621.939 Ha", "14,76%", "+316.766 Ha", "+6,67%"],
        ["Sangat Tinggi", "4.2–5.0", "2.870.741 Ha", "65,54%", "1.941.460 Ha", "46,07%", "+929.281 Ha", "+19,47%"],
        ["TOTAL DARATAN", "-", "4.380.394 Ha", "100,00%", "4.213.812 Ha", "100,00%", "+166.582 Ha", "-"]
    ]
    add_matrix_table(slide7, common_headers, pku_data)
    add_insight_box(slide7, "Analisis Spasial PKU", [
        ("Keunggulan Deteksi Kanopi", "Raster 0.5m mencatat kelas Sangat Tinggi sebesar 65,54% (2,87 juta Ha), jauh melampaui GDB (46,07%)."),
        ("Efektivitas Serapan Udara", "Kerapatan kanopi vegetasi tropika di lereng Bukit Barisan dan Mentawai secara kontinu memproduksi oksigen dan meredam pulau panas."),
        ("Area Terpolusi Mikro", "Lahan terbangun dan jalan aspal kota (Padang, Bukittinggi) terdeteksi presisi di kelas Sangat Rendah & Rendah (~51.000 Ha).")
    ])

    # =========================================================================
    # SLIDE 8: PHK (HABITAT & KEHATI)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Matriks Perbandingan: PHK (Habitat & Keanekaragaman Hayati)")
    
    phk_headers = ["Kelas JLH", "Skor", "Rast 0.5m (Alami)", "% Rast", "GDB (Alami)", "% Alami", "GDB (dengan KK)", "% KK"]
    phk_data = [
        ["Sangat Rendah", "1.0–1.8", "184.683 Ha", "4,22%", "116.052 Ha", "2,75%", "87.636 Ha", "2,08%"],
        ["Rendah", "1.8–2.6", "769.738 Ha", "17,57%", "1.411.072 Ha", "33,49%", "1.397.909 Ha", "33,17%"],
        ["Sedang", "2.6–3.4", "700.684 Ha", "16,00%", "1.069.286 Ha", "25,38%", "806.200 Ha", "19,13%"],
        ["Tinggi", "3.4–4.2", "1.363.216 Ha", "31,12%", "1.401.481 Ha", "33,26%", "342.660 Ha", "8,13%"],
        ["Sangat Tinggi", "4.2–5.0", "1.362.073 Ha", "31,09%", "215.922 Ha", "5,12%", "1.579.407 Ha", "37,48%"],
        ["TOTAL DARATAN", "-", "4.380.394 Ha", "100,00%", "4.213.812 Ha", "100,00%", "4.213.812 Ha", "100,00%"]
    ]
    add_matrix_table(slide8, phk_headers, phk_data)
    add_insight_box(slide8, "Analisis Kriteria Khusus", [
        ("Efek Mandat Konservasi", "Penerapan poligon KK.gpkg (Kriteria Khusus 1) melonjakkan kelas Sangat Tinggi pada GDB dari 5,12% menjadi 37,48% (+1,36 juta Ha)."),
        ("Konservasi Mutlak", "Kawasan TNKS, Cagar Alam Batang Palupuh, dan Suaka Margasatwa dipaksa bernilai 5.0 sebagai koridor inti satwa kunci."),
        ("Konsistensi Raster & Vektor", "Raster PHK Alami telah menghasilkan 31,09% kelas Sangat Tinggi bahkan sebelum diterapkan KK, menunjukkan akurasi tinggi klasifikasi hutan primer.")
    ])

    # =========================================================================
    # SLIDE 9: PPK (PENYIMPAN KARBON)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Matriks Perbandingan: PPK (Penyerap & Penyimpan Karbon / Carbon Storage)")
    ppk_data = [
        ["Sangat Rendah", "1.0–1.8", "134.381 Ha", "3,07%", "144.213 Ha", "3,42%", "-9.832 Ha", "-0,35%"],
        ["Rendah", "1.8–2.6", "639.492 Ha", "14,60%", "2.116.121 Ha", "50,22%", "-1.476.629 Ha", "-35,62%"],
        ["Sedang", "2.6–3.4", "930.350 Ha", "21,24%", "101.914 Ha", "2,42%", "+828.436 Ha", "+18,82%"],
        ["Tinggi", "3.4–4.2", "170.338 Ha", "3,89%", "640.778 Ha", "15,21%", "-470.440 Ha", "-11,32%"],
        ["Sangat Tinggi", "4.2–5.0", "2.505.834 Ha", "57,21%", "1.210.786 Ha", "28,73%", "+1.295.048 Ha", "+28,48%"],
        ["TOTAL DARATAN", "-", "4.380.394 Ha", "100,00%", "4.213.812 Ha", "100,00%", "+166.582 Ha", "-"]
    ]
    add_matrix_table(slide9, common_headers, ppk_data)
    add_insight_box(slide9, "Analisis Spasial PPK", [
        ("Potensi Cadangan Karbon", "Raster 0.5m mengidentifikasi 2,50 juta Ha (57,21%) stok karbon Sangat Tinggi, berbanding 1,21 juta Ha (28,73%) pada GDB."),
        ("Dukungan FOLU Net Sink", "Data 0.5m membuktikan biomassa tegakan hutan dan tanah gambut pesisir barat Sumbar memiliki cadangan simpanan karbon yang jauh lebih besar dari estimasi makro."),
        ("Penyusutan Kelas Rendah", "Kelas Rendah di GDB (50,22%) terbukti merupakan dampak generalisasi kebun karet/sawit tua yang sebenarnya menyimpan biomassa signifikan.")
    ])

    # =========================================================================
    # SLIDE 10: REKOMENDASI PENGGUNAAN
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Sintesis & Panduan Kebijakan Penggunaan Data")
    
    use_cards = [
        ("KAPAN MEMAKAI RASTER 0.5 METER?", [
            ("Analisis Skala Tapak", "Penyusunan dokumen AMDAL, UKL-UPL, dan perizinan berusaha PBPH/Persetujuan Lingkungan."),
            ("Mitigasi Bencana & Relokasi", "Pemilihan tapak hunian tetap/sementara yang aman bencana dan memiliki daya dukung air tinggi."),
            ("Kajian D3TLH Mikro", "Penentuan batas sempadan sungai 5-10m, koridor jelajah satwa liar, dan perlindungan mata air.")
        ], Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4)),
        ("KAPAN MEMAKAI VEKTOR GDB?", [
            ("Laporan Makro RPPLH", "Penyusunan buku profil lingkungan hidup tingkat provinsi/kabupaten dan neraca SDA tahunan."),
            ("Tabulasi Administrasi", "Query cepat luasan JLH agregat per kecamatan, kabupaten/kota, atau Daerah Aliran Sungai (DAS)."),
            ("Penyelarasan RTRW Provinsi", "Integrasi pola ruang makro dan evaluasi Kajian Lingkungan Hidup Strategis (KLHS) tingkat provinsi.")
        ], Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.4))
    ]
    for c_title, c_bullets, c_l, c_t, c_w, c_h in use_cards:
        add_insight_box(slide10, c_title, c_bullets, left=c_l, top=c_t, width=c_w, height=c_h)

    # Save
    out_pptx = "Matriks_Perbandingan_JLH_Sumbar_2025.pptx"
    prs.save(out_pptx)
    print(f"Presentasi PowerPoint berhasil dibuat: {out_pptx}")

if __name__ == "__main__":
    create_presentation()
