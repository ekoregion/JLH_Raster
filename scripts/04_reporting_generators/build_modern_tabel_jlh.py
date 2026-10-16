import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 1. Inisialisasi Presentasi Widescreen 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Skema Warna Modern
COLOR_DARK_BG = RGBColor(15, 23, 42)      # Slate 900
COLOR_TEXT_DARK = RGBColor(15, 23, 42)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)
COLOR_WHITE = RGBColor(255, 255, 255)

COLOR_RASTER = RGBColor(2, 132, 199)      # Biru Sky 600
COLOR_VEKTOR = RGBColor(13, 148, 136)     # Toska Teal 600
COLOR_SELISIH = RGBColor(51, 65, 85)      # Slate 700
COLOR_KK = RGBColor(99, 102, 241)         # Indigo 500

COLOR_POS = RGBColor(21, 128, 61)         # Hijau Positif
COLOR_NEG = RGBColor(220, 38, 38)         # Merah Negatif

# Warna badge latar belakang kategori JLH
BADGE_COLORS = {
    "Sangat Rendah": (RGBColor(254, 226, 226), RGBColor(185, 28, 28)),
    "Rendah": (RGBColor(254, 243, 199), RGBColor(180, 83, 9)),
    "Sedang": (RGBColor(254, 249, 195), RGBColor(133, 77, 14)),
    "Tinggi": (RGBColor(220, 252, 231), RGBColor(21, 128, 61)),
    "Sangat Tinggi": (RGBColor(209, 250, 229), RGBColor(4, 120, 87)),
}

# --- SLIDE 1: COVER ---
slide1 = prs.slides.add_slide(blank_layout)
bg_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg_box.fill.solid()
bg_box.fill.fore_color.rgb = COLOR_DARK_BG
bg_box.line.fill.background()

tb = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
tf = tb.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "KEMENTERIAN LINGKUNGAN HIDUP / BPLH • DIREKTORAT PENYELENGGARAAN SDA BERKELANJUTAN"
p0.font.size = Pt(13)
p0.font.bold = True
p0.font.color.rgb = RGBColor(56, 189, 248)

p1 = tf.add_paragraph()
p1.text = "Matriks Komparasi Data Raster vs GDB Vektor"
p1.font.size = Pt(36)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.space_before = Pt(14)

p2 = tf.add_paragraph()
p2.text = "Visualisasi dan evaluasi spasial Kelas Jasa Lingkungan Hidup (JLH) pada seluruh tema fungsi ekosistem daratan."
p2.font.size = Pt(16)
p2.font.color.rgb = RGBColor(148, 163, 184)
p2.space_before = Pt(12)

# --- DAFTAR DATA SHEET ---
data_sheets = [
    {
        "code": "PGA",
        "title": "Matriks Perbandingan: PGA (Pengaturan Air / Water Regulation)",
        "scope": "Water Regulation",
        "rows": [
            ("Sangat Rendah", "1.0–1.8", "29.073", "0,66%", "8.993", "0,21%", "+20.080 Ha", "+0,45%"),
            ("Rendah", "1.8–2.6", "357.572", "8,16%", "309.061", "7,33%", "+48.511 Ha", "+0,83%"),
            ("Sedang", "2.6–3.4", "1.149.438", "26,24%", "1.751.883", "41,57%", "-602.445 Ha", "-15,33%"),
            ("Tinggi", "3.4–4.2", "733.377", "16,74%", "716.062", "16,99%", "+17.315 Ha", "-0,25%"),
            ("Sangat Tinggi", "4.2–5.0", "2.110.935", "48,19%", "1.427.814", "33,88%", "+683.121 Ha", "+14,31%"),
        ],
        "total": ("TOTAL DARATAN", "-", "4.380.394", "100,00%", "4.213.812", "100,00%", "+166.582 Ha", "-")
    },
    {
        "code": "PYA",
        "title": "Matriks Perbandingan: PYA (Penyedia Air Bersih / Water Provisioning)",
        "scope": "Water Provisioning",
        "rows": [
            ("Sangat Rendah", "1.0–1.8", "92.956", "2,12%", "46.574", "1,11%", "+46.382 Ha", "+1,01%"),
            ("Rendah", "1.8–2.6", "891.305", "20,35%", "2.364.620", "56,12%", "-1.473.315 Ha", "-35,77%"),
            ("Sedang", "2.6–3.4", "3.113.503", "71,08%", "1.749.860", "41,53%", "+1.363.643 Ha", "+29,55%"),
            ("Tinggi", "3.4–4.2", "240.258", "5,48%", "17.717", "0,42%", "+222.541 Ha", "+5,06%"),
            ("Sangat Tinggi", "4.2–5.0", "42.372", "0,97%", "35.041", "0,83%", "+7.331 Ha", "+0,14%"),
        ],
        "total": ("TOTAL DARATAN", "-", "4.380.394", "100,00%", "4.213.812", "100,00%", "+166.582 Ha", "-")
    },
    {
        "code": "PGN",
        "title": "Matriks Perbandingan: PGN (Penyedia Pangan / Food Provisioning)",
        "scope": "Food Provisioning",
        "rows": [
            ("Sangat Rendah", "1.0–1.8", "38.670", "0,88%", "26.376", "0,63%", "+12.294 Ha", "+0,25%"),
            ("Rendah", "1.8–2.6", "307.563", "7,02%", "113.411", "2,69%", "+194.152 Ha", "+4,33%"),
            ("Sedang", "2.6–3.4", "671.672", "15,33%", "2.200.546", "52,22%", "-1.528.874 Ha", "-36,89%"),
            ("Tinggi", "3.4–4.2", "3.191.533", "72,86%", "1.641.998", "38,97%", "+1.549.535 Ha", "+33,89%"),
            ("Sangat Tinggi", "4.2–5.0", "170.956", "3,90%", "231.482", "5,49%", "-60.526 Ha", "-1,59%"),
        ],
        "total": ("TOTAL DARATAN", "-", "4.380.394", "100,00%", "4.213.812", "100,00%", "+166.582 Ha", "-")
    },
    {
        "code": "PKU",
        "title": "Matriks Perbandingan: PKU (Pengatur Kualitas Udara / Air Quality)",
        "scope": "Air Quality",
        "rows": [
            ("Sangat Rendah", "1.0–1.8", "6.840", "0,16%", "0", "0,00%", "+6.840 Ha", "+0,16%"),
            ("Rendah", "1.8–2.6", "44.263", "1,01%", "22.878", "0,54%", "+21.385 Ha", "+0,47%"),
            ("Sedang", "2.6–3.4", "519.845", "11,87%", "1.627.534", "38,62%", "-1.107.689 Ha", "-26,75%"),
            ("Tinggi", "3.4–4.2", "938.705", "21,43%", "621.939", "14,76%", "+316.766 Ha", "+6,67%"),
            ("Sangat Tinggi", "4.2–5.0", "2.870.741", "65,54%", "1.941.460", "46,07%", "+929.281 Ha", "+19,47%"),
        ],
        "total": ("TOTAL DARATAN", "-", "4.380.394", "100,00%", "4.213.812", "100,00%", "+166.582 Ha", "-")
    },
    {
        "code": "PHK",
        "title": "Matriks Perbandingan: PHK (Habitat & Keanekaragaman Hayati)",
        "scope": "Biodiversity & Habitat",
        "headers": ["Kelas JLH", "Skor", "Rast 0.5m (Alami)", "% Rast", "GDB (Alami)", "% Alami", "GDB (dengan KK)", "% KK"],
        "rows": [
            ("Sangat Rendah", "1.0–1.8", "184.683", "4,22%", "116.052", "2,75%", "87.636", "2,08%"),
            ("Rendah", "1.8–2.6", "769.738", "17,57%", "1.411.072", "33,49%", "1.397.909", "33,17%"),
            ("Sedang", "2.6–3.4", "700.684", "16,00%", "1.069.286", "25,38%", "806.200", "19,13%"),
            ("Tinggi", "3.4–4.2", "1.363.216", "31,12%", "1.401.481", "33,26%", "342.660", "8,13%"),
            ("Sangat Tinggi", "4.2–5.0", "1.362.073", "31,09%", "215.922", "5,12%", "1.579.407", "37,48%"),
        ],
        "total": ("TOTAL DARATAN", "-", "4.380.394", "100,00%", "4.213.812", "100,00%", "4.213.812", "100,00%")
    },
    {
        "code": "PPK",
        "title": "Matriks Perbandingan: PPK (Penyerap & Penyimpan Karbon)",
        "scope": "Carbon Storage",
        "rows": [
            ("Sangat Rendah", "1.0–1.8", "134.381", "3,07%", "144.213", "3,42%", "-9.832 Ha", "-0,35%"),
            ("Rendah", "1.8–2.6", "639.492", "14,60%", "2.116.121", "50,22%", "-1.476.629 Ha", "-35,62%"),
            ("Sedang", "2.6–3.4", "930.350", "21,24%", "101.914", "2,42%", "+828.436 Ha", "+18,82%"),
            ("Tinggi", "3.4–4.2", "170.338", "3,89%", "640.778", "15,21%", "-470.440 Ha", "-11,32%"),
            ("Sangat Tinggi", "4.2–5.0", "2.505.834", "57,21%", "1.210.786", "28,73%", "+1.295.048 Ha", "+28,48%"),
        ],
        "total": ("TOTAL DARATAN", "-", "4.380.394", "100,00%", "4.213.812", "100,00%", "+166.582 Ha", "-")
    }
]

# --- LOOP BUILD TABLE SLIDES ---
for item in data_sheets:
    slide = prs.slides.add_slide(blank_layout)
    
    # 1. Header slide
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    htf = header_box.text_frame
    htf.word_wrap = True
    
    hp0 = htf.paragraphs[0]
    hp0.text = "KEMENTERIAN LINGKUNGAN HIDUP / BPLH • DIREKTORAT PENYELENGGARAAN SDA BERKELANJUTAN"
    hp0.font.size = Pt(9.5)
    hp0.font.bold = True
    hp0.font.color.rgb = COLOR_VEKTOR
    
    hp1 = htf.add_paragraph()
    hp1.text = item["title"]
    hp1.font.size = Pt(19)
    hp1.font.bold = True
    hp1.font.color.rgb = COLOR_TEXT_DARK
    
    # 2. Tabel
    rows_cnt = 2 + len(item["rows"]) + 1  # Group Header + Sub Header + Data + Total
    cols_cnt = 8
    tbl_shape = slide.shapes.add_table(rows_cnt, cols_cnt, Inches(0.8), Inches(1.5), Inches(11.73), Inches(4.8))
    tbl = tbl_shape.table
    
    # Atur lebar kolom
    tbl.columns[0].width = Inches(1.8)  # Kelas JLH
    tbl.columns[1].width = Inches(1.1)  # Skor
    tbl.columns[2].width = Inches(1.5)  # Luas Rast
    tbl.columns[3].width = Inches(1.3)  # % Rast
    tbl.columns[4].width = Inches(1.5)  # Luas Vekt
    tbl.columns[5].width = Inches(1.3)  # % Vekt
    tbl.columns[6].width = Inches(1.6)  # Selisih Luas
    tbl.columns[7].width = Inches(1.63) # Selisih %

    # Group Header (Row 0)
    cell_klas = tbl.cell(0, 0)
    cell_klas.merge(tbl.cell(0, 1))
    cell_klas.text = "KLASIFIKASI"
    cell_klas.fill.solid()
    cell_klas.fill.fore_color.rgb = RGBColor(71, 85, 105)

    cell_rast = tbl.cell(0, 2)
    cell_rast.merge(tbl.cell(0, 3))
    cell_rast.text = "RASTER (0.5m)"
    cell_rast.fill.solid()
    cell_rast.fill.fore_color.rgb = COLOR_RASTER

    cell_vekt = tbl.cell(0, 4)
    cell_vekt.merge(tbl.cell(0, 5))
    cell_vekt.text = "GDB VEKTOR"
    cell_vekt.fill.solid()
    cell_vekt.fill.fore_color.rgb = COLOR_VEKTOR

    cell_sel = tbl.cell(0, 6)
    cell_sel.merge(tbl.cell(0, 7))
    if item["code"] == "PHK":
        cell_sel.text = "GDB (DENGAN KK)"
        cell_sel.fill.solid()
        cell_sel.fill.fore_color.rgb = COLOR_KK
    else:
        cell_sel.text = "SELISIH (RASTER - VEKTOR)"
        cell_sel.fill.solid()
        cell_sel.fill.fore_color.rgb = COLOR_SELISIH

    for c in [cell_klas, cell_rast, cell_vekt, cell_sel]:
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_WHITE

    # Sub Header (Row 1)
    if "headers" in item:
        subheaders = item["headers"]
    else:
        subheaders = ["Kelas JLH", "Skor", "Luas (Ha)", "% Raster", "Luas (Ha)", "% Vektor", "Selisih Luas", "Selisih %"]
    
    for c_idx, sh in enumerate(subheaders):
        cell = tbl.cell(1, c_idx)
        cell.text = sh
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.alignment = PP_ALIGN.CENTER if c_idx in [0, 1] else PP_ALIGN.RIGHT

    # Baris Data (Row 2 .. 6)
    for r_idx, row_data in enumerate(item["rows"]):
        for c_idx, val in enumerate(row_data):
            cell = tbl.cell(2 + r_idx, c_idx)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10.5)

            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if r_idx % 2 == 0 else RGBColor(248, 250, 252)

            if c_idx == 0:
                p.alignment = PP_ALIGN.CENTER
                p.font.bold = True
                bg_col, txt_col = BADGE_COLORS.get(val, (RGBColor(241, 245, 249), COLOR_TEXT_DARK))
                cell.fill.fore_color.rgb = bg_col
                p.font.color.rgb = txt_col
            elif c_idx == 1:
                p.alignment = PP_ALIGN.CENTER
                p.font.color.rgb = COLOR_TEXT_MUTED
            elif c_idx in [2, 4]:
                p.alignment = PP_ALIGN.RIGHT
                p.font.color.rgb = COLOR_TEXT_DARK
                p.font.bold = True
            elif c_idx in [3, 5]:
                p.alignment = PP_ALIGN.RIGHT
                p.font.color.rgb = COLOR_TEXT_DARK
            elif c_idx in [6, 7]:
                p.alignment = PP_ALIGN.RIGHT
                p.font.bold = True
                if "-" in val and val != "-":
                    p.font.color.rgb = COLOR_NEG
                elif "+" in val:
                    p.font.color.rgb = COLOR_POS
                else:
                    p.font.color.rgb = COLOR_TEXT_DARK

    # Baris Total (Row Terakhir)
    tot_row_idx = 2 + len(item["rows"])
    for c_idx, val in enumerate(item["total"]):
        cell = tbl.cell(tot_row_idx, c_idx)
        cell.text = val
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_DARK_BG
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER if c_idx in [0, 1] else PP_ALIGN.RIGHT

    # Footnote
    fn_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.4))
    fn_p = fn_box.text_frame.paragraphs[0]
    fn_p.text = "*Ha = Hektar | JLH = Jasa Lingkungan Hidup | Desain Presentasi Standar KLH-BPLH 2026"
    fn_p.font.size = Pt(9)
    fn_p.font.color.rgb = COLOR_TEXT_MUTED

# Simpan ke kedua file
prs.save("tabel JLH.pptx")
prs.save("Matriks_Perbandingan_JLH_Modern.pptx")
print("Berhasil memperbarui 'tabel JLH.pptx' dan membuat 'Matriks_Perbandingan_JLH_Modern.pptx'")
